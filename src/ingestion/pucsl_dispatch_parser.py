"""
Source-specific extractor for PUCSL 15-minute generation dispatch XLSX workbooks.
Extracts 96 sub-daily discrete intervals per station per day, preserving sheet references,
cell coordinates, and distinguishing plant shutdown (0.0 MW) from unrecorded missing intervals (None).
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import openpyxl

COLOMBO_TZ = ZoneInfo("Asia/Colombo")


@dataclass(frozen=True)
class DispatchIntervalRecord:
    """Sub-daily 15-minute dispatch observation extracted from PUCSL dispatch sheet."""

    observation_date: date
    interval_slot_number: int
    observation_timestamp: datetime
    raw_entity_name: str
    power_mw: float | None
    is_missing: bool
    reported_unit_code: str = "MW"
    sheet_name: str = ""
    row_index: int = 0
    column_index: int = 0


class PucslDispatchParser:
    """Parses PUCSL 15-minute interval dispatch workbooks."""

    def get_sheet_names(self, file_path: Path) -> list[str]:
        """Return list of sheet names in the dispatch workbook."""
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        try:
            return [s.strip() for s in wb.sheetnames]
        finally:
            wb.close()

    @staticmethod
    def _parse_sheet_date(sheet_name: str, row0_val: Any) -> date:
        """Infer calendar date from row 0 metadata or sheet name (e.g. '01.03')."""
        if isinstance(row0_val, datetime):
            return row0_val.date()
        if isinstance(row0_val, date):
            return row0_val

        # Fallback to sheet name matching DD.MM
        m = re.search(r"(\d{1,2})\.(\d{1,2})", sheet_name)
        if m:
            day = int(m.group(1))
            month = int(m.group(2))
            return date(2024, month, day)

        raise ValueError(
            f"Cannot resolve calendar date from sheet '{sheet_name}' and row0: {row0_val}"
        )

    def parse_sheet(self, file_path: Path, sheet_name: str) -> list[DispatchIntervalRecord]:
        """Parse all 96 interval slots for all stations on the specified daily sheet."""
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        records: list[DispatchIntervalRecord] = []

        try:
            # Handle potential whitespace in sheet names
            matched_sheet = None
            for name in wb.sheetnames:
                if name.strip() == sheet_name.strip():
                    matched_sheet = name
                    break

            if not matched_sheet:
                raise ValueError(f"Sheet '{sheet_name}' not found in workbook {file_path.name}")

            ws = wb[matched_sheet]
            all_rows = list(ws.iter_rows(values_only=True))
            if len(all_rows) < 3:
                return []

            # Row 0: Header with date
            row0 = all_rows[0]
            sheet_date = self._parse_sheet_date(sheet_name, row0[8] if len(row0) > 8 else None)

            # Row 1: Time labels (e.g. Time, 0:15, 0:30, ..., 24:00)
            all_rows[1]
            time_slots: list[tuple[int, time]] = []
            for col_idx in range(1, 97):
                total_min = col_idx * 15
                if total_min == 1440:
                    # 24:00 is represented as end of day
                    slot_time = time(23, 59, 59)
                else:
                    slot_time = time(total_min // 60, total_min % 60, 0)
                time_slots.append((col_idx, slot_time))

            # Rows 2+: Power stations and aggregates
            for r_idx, row in enumerate(all_rows[2:], start=2):
                if not row or row[0] is None:
                    continue

                raw_station = str(row[0]).strip()
                if not raw_station or raw_station.upper() in ("TIME", "ACTUAL SYSTEM DISPATCH"):
                    continue

                # Stop reading once we reach the summary block starting with Total MW
                if raw_station.upper().startswith("TOTAL MW") or raw_station.upper() == "TOTAL":
                    break

                for col_idx, (slot_num, slot_time) in enumerate(time_slots, start=1):
                    val = row[col_idx] if col_idx < len(row) else None
                    power_mw: float | None = None
                    is_missing = False

                    if val is None or str(val).strip() in ("", "-", "N/A", "null"):
                        power_mw = None
                        is_missing = True
                    else:
                        try:
                            power_mw = float(val)
                        except (ValueError, TypeError):
                            power_mw = None
                            is_missing = True

                    # Construct timezone-aware SLST timestamp
                    if slot_num == 96:
                        # 24:00 corresponds to next day 00:00:00 SLST
                        next_day = sheet_date + timedelta(days=1)
                        ts = datetime(
                            next_day.year, next_day.month, next_day.day, 0, 0, 0, tzinfo=COLOMBO_TZ
                        )
                    else:
                        ts = datetime(
                            sheet_date.year,
                            sheet_date.month,
                            sheet_date.day,
                            slot_time.hour,
                            slot_time.minute,
                            0,
                            tzinfo=COLOMBO_TZ,
                        )

                    records.append(
                        DispatchIntervalRecord(
                            observation_date=sheet_date,
                            interval_slot_number=slot_num,
                            observation_timestamp=ts,
                            raw_entity_name=raw_station,
                            power_mw=power_mw,
                            is_missing=is_missing,
                            reported_unit_code="MW",
                            sheet_name=matched_sheet,
                            row_index=r_idx,
                            column_index=col_idx,
                        )
                    )

            return records
        finally:
            wb.close()

    def parse_all_sheets(
        self, file_path: Path, max_sheets: int | None = None
    ) -> list[DispatchIntervalRecord]:
        """Parse all daily sheets in the workbook (optionally limited to max_sheets)."""
        sheet_names = self.get_sheet_names(file_path)
        if max_sheets:
            sheet_names = sheet_names[:max_sheets]

        all_records: list[DispatchIntervalRecord] = []
        for s in sheet_names:
            all_records.extend(self.parse_sheet(file_path, s))
        return all_records
